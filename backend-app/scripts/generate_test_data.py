from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
import xlrd

ROOT = Path(r'D:\Obsidian\web\llm-wiki\raw\inbox\projects\DermaIntegrate\backend-app')
TARGET = ROOT / '测试数据'
FONT_PATH = Path(r'C:\Windows\Fonts\msyh.ttc')

KB_SPECS = [
    {
        'folder': '医生张晨个人库-痤疮与炎症管理',
        'theme': '围绕痤疮、毛囊炎、玫瑰痤疮以及常见炎症性皮肤问题的门诊记录、复盘思路与随访提醒。',
        'focus': '强调病程时间线、诱因识别、局部处理、口服用药观察与复诊节点管理。',
    },
    {
        'folder': '医生李雯个人库-皮肤镜随访复盘',
        'theme': '围绕皮肤镜随访、色素性病变观察、术后复盘与沟通话术整理的个人知识库。',
        'focus': '强调图像对照、风险特征记录、复查提醒、沟通边界与证据不足时的保守表达。',
    },
    {
        'folder': '管理员个人库-知识治理与标准',
        'theme': '围绕知识库命名规范、文档审核规则、版本治理、权限管理与数据质量检查的管理员个人资料。',
        'focus': '强调结构统一、命名一致、敏感信息脱敏、上传审查与维护清单。',
    },
    {
        'folder': '皮肤科科室库-MDT与病理协作',
        'theme': '围绕科室内 MDT 会诊、病理回传、临床协作、转诊建议与多学科沟通流程的科室知识库。',
        'focus': '强调会诊记录、责任分工、行动项追踪、病理字段复核与临床路径衔接。',
    },
    {
        'folder': '公共知识库-患者宣教与早筛',
        'theme': '围绕患者宣教、早筛提示、复诊提醒、家庭观察与就医建议的公共知识库。',
        'focus': '强调通俗表达、风险信号、拍照留存、防晒建议与什么时候必须就医。',
    },
]

FILES = [
    '01-主题说明.md',
    '02-知识摘要.txt',
    '03-结构化清单.csv',
    '04-结构化清单.xlsx',
    '05-结构化清单.xls',
    '06-详细说明.docx',
    '07-详细说明.pdf',
]

pdfmetrics.registerFont(TTFont('MSYH', str(FONT_PATH)))
styles = getSampleStyleSheet()
BODY = ParagraphStyle('BodyCN', parent=styles['BodyText'], fontName='MSYH', fontSize=11, leading=16, spaceAfter=6)
TITLE = ParagraphStyle('TitleCN', parent=styles['Title'], fontName='MSYH', fontSize=16, leading=20, spaceAfter=10)
H2 = ParagraphStyle('H2CN', parent=styles['Heading2'], fontName='MSYH', fontSize=12, leading=16, spaceAfter=8)


def long_text(spec: dict, label: str) -> str:
    return (
        f"{spec['folder']}属于中文测试知识库，主题是：{spec['theme']}"
        f"本库用于验证上传、解析、切分、检索与预览流程，文档内容专门围绕真实中文医学场景组织，"
        f"避免空泛描述。{label}部分重点说明：{spec['focus']}"
        f"在实际使用中，建议把每条记录写成完整句子，尽量包含背景、症状、时间、处理方式和下一步建议。"
        f"这样既便于人工阅读，也便于向量化和关键词检索命中。"
        f"本测试资料统一采用 UTF-8 编码，不使用 GBK，不留乱码占位，也不写模糊的替代符号。"
        f"如果后续需要扩充内容，可以继续沿着这个主题增加病例讨论、护理提醒、宣教示例与版本对照说明。"
    )


def markdown_text(spec: dict) -> str:
    return (
        f"# {spec['folder']}\n\n"
        f"{long_text(spec, 'Markdown 文档')}\n\n"
        f"## 使用说明\n"
        f"1. 该知识库文件用于测试中文内容解析。\n"
        f"2. 所有正文都要求可直接阅读，不使用乱码占位。\n"
        f"3. 结构上保留标题、段落与列表，方便检索系统切分。\n"
        f"4. 未来可以继续补充版本记录和修订说明。\n"
    )


def plain_text(spec: dict) -> str:
    return (
        f"{spec['folder']}\n\n"
        f"{long_text(spec, '纯文本说明')}\n\n"
        f"补充说明：本文件用于检查系统对 UTF-8 中文纯文本的读取能力，内容保持连续、完整、可分段。"
        f"在测试时可以直接把它当作门诊摘要、内部说明或知识备注来解析。"
    )


def rows(spec: dict):
    data = []
    for i in range(1, 6):
        data.append([
            i,
            f"场景{i}",
            f"{spec['folder']}的第{i}条结构化记录强调：{spec['focus']}。该条内容用于测试表格型文档的逐行提取能力，要求既包含字段名，也包含完整语义。",
            f"{spec['theme']} 该说明用于补充行级上下文，便于检索时把字段和值一起命中，并能覆盖更长的中文叙述。",
        ])
    return data


def write_csv_file(path: Path, spec: dict):
    with path.open('w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['序号', '场景', '记录重点', '补充说明'])
        writer.writerows(rows(spec))


def write_xlsx_file(path: Path, spec: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = '测试数据'
    ws.append(['序号', '场景', '记录重点', '补充说明'])
    for row in rows(spec):
        ws.append(row)
    wb.save(path)


def write_xls_file(path: Path, spec: dict):
    import xlwt

    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('测试数据')
    headers = ['序号', '场景', '记录重点', '补充说明']
    for c, h in enumerate(headers):
        ws.write(0, c, h)
    for r, row in enumerate(rows(spec), start=1):
        for c, value in enumerate(row):
            ws.write(r, c, value)
    wb.save(str(path))


def write_docx_file(path: Path, spec: dict):
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading(spec['folder'], level=1)
    for label in ['Word 文档的第一段', 'Word 文档的第二段，补充门诊背景、记录方式与复盘要求', 'Word 文档的第三段，强调测试时需要检查段落顺序和中文编码']:
        doc.add_paragraph(long_text(spec, label))
    doc.save(path)


def write_pdf_file(path: Path, spec: dict):
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    story = [Paragraph(spec['folder'], TITLE), Spacer(1, 4)]
    for heading, label in [('测试背景', 'PDF 第一段，用于说明主题与验证点'), ('记录要求', 'PDF 第二段，说明中文段落与分行规范'), ('补充说明', 'PDF 第三段，说明后续扩展和版本维护方式')]:
        story.append(Paragraph(heading, H2))
        story.append(Paragraph(long_text(spec, label), BODY))
    doc.build(story)


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return '\n'.join(p.text for p in doc.paragraphs)


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return '\n'.join(page.extract_text() or '' for page in reader.pages)


def extract_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True)
    ws = wb.active
    return '\n'.join(' '.join('' if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))


def extract_xls(path: Path) -> str:
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    lines = []
    for r in range(sheet.nrows):
        lines.append(' '.join(str(sheet.cell_value(r, c)) for c in range(sheet.ncols)))
    return '\n'.join(lines)


if TARGET.exists():
    shutil.rmtree(TARGET)
TARGET.mkdir(parents=True, exist_ok=True)

manifest = []
for spec in KB_SPECS:
    folder = TARGET / spec['folder']
    folder.mkdir(parents=True, exist_ok=True)

    (folder / '01-主题说明.md').write_text(markdown_text(spec), encoding='utf-8')
    (folder / '02-知识摘要.txt').write_text(plain_text(spec), encoding='utf-8')
    write_csv_file(folder / '03-结构化清单.csv', spec)
    write_xlsx_file(folder / '04-结构化清单.xlsx', spec)
    write_xls_file(folder / '05-结构化清单.xls', spec)
    write_docx_file(folder / '06-详细说明.docx', spec)
    write_pdf_file(folder / '07-详细说明.pdf', spec)

    manifest.append({'folder': spec['folder'], 'files': FILES})

(TARGET / 'manifest.json').write_text(
    json.dumps({'root': str(TARGET.relative_to(ROOT)), 'kb_count': len(KB_SPECS), 'items': manifest}, ensure_ascii=False, indent=2),
    encoding='utf-8'
)
(TARGET / 'README.md').write_text('# 测试数据\n\n本目录包含 5 套中文测试知识库，每套知识库都包含系统支持的常见文件类型。\n', encoding='utf-8')

verification = []
for spec in KB_SPECS:
    folder = TARGET / spec['folder']
    texts = {
        '01-主题说明.md': (folder / '01-主题说明.md').read_text(encoding='utf-8'),
        '02-知识摘要.txt': (folder / '02-知识摘要.txt').read_text(encoding='utf-8'),
        '03-结构化清单.csv': (folder / '03-结构化清单.csv').read_text(encoding='utf-8'),
        '04-结构化清单.xlsx': extract_xlsx(folder / '04-结构化清单.xlsx'),
        '05-结构化清单.xls': extract_xls(folder / '05-结构化清单.xls'),
        '06-详细说明.docx': extract_docx(folder / '06-详细说明.docx'),
        '07-详细说明.pdf': extract_pdf(folder / '07-详细说明.pdf'),
    }
    sizes = {}
    for name, text in texts.items():
        if '??' in text:
            raise RuntimeError(f'Found placeholder ?? in {folder / name}')
        if '�' in text:
            raise RuntimeError(f'Found replacement character in {folder / name}')
        cleaned = text.replace('\r', '').replace('\n', '').strip()
        if len(cleaned) < 200:
            raise RuntimeError(f'Text too short in {folder / name}: {len(cleaned)} chars')
        sizes[name] = len(cleaned)
    verification.append({'folder': spec['folder'], 'sizes': sizes})

print(json.dumps({'status': 'ok', 'created_root': str(TARGET), 'verification': verification}, ensure_ascii=False, indent=2))
