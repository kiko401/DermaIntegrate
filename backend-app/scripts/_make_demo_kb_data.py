from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

ROOT = Path.cwd() / 'backend-app'
TARGET = ROOT / '\u6d4b\u8bd5\u6570\u636e'

MD_NAME = '\u0030\u0031\u002d\u4e3b\u9898\u603b\u89c8\u002e\u006d\u0064'
TXT_NAME = '\u0030\u0032\u002d\u95e8\u8bca\u5b66\u4e60\u6458\u8981\u002e\u0074\u0078\u0074'
CSV_NAME = '\u0030\u0033\u002d\u7ed3\u6784\u5316\u8981\u70b9\u002e\u0063\u0073\u0076'
XLSX_NAME = '\u0030\u0034\u002d\u7ed3\u6784\u5316\u8981\u70b9\u002e\u0078\u006c\u0073\u0078'
XLS_NAME = '\u0030\u0035\u002d\u7ed3\u6784\u5316\u8981\u70b9\u002e\u0078\u006c\u0073'
DOCX_NAME = '\u0030\u0036\u002d\u533b\u751f\u6ce8\u610f\u4e8b\u9879\u002e\u0064\u006f\u0063\u0078'
PDF_NAME = '\u0030\u0037\u002d\u5b66\u4e60\u8d44\u6599\u6c47\u7f16\u002e\u0070\u0064\u0066'

SUITES = [
    {
        'folder': '\u6f14\u793a\u533b\u751f\u5e93-\u9ed1\u8272\u7d20\u7624\u89c4\u8303\u5b66\u4e60\u8d44\u6599',
        'title': '\u9ed1\u8272\u7d20\u7624\u89c4\u8303\u5b66\u4e60\u8d44\u6599',
        'theme': '\u56f4\u7ed5\u9ed1\u8272\u7d20\u7624\u65e9\u671f\u8bc6\u522b\u3001\u95e8\u8bca\u6c9f\u901a\u3001\u75c5\u7406\u534f\u4f5c\u4e0e\u968f\u8bbf\u7ba1\u7406\u6574\u7406\u7684\u5b66\u4e60\u8d44\u6599\u3002',
        'focus': '\u5f3a\u8c03\u95e8\u8bca\u533b\u751f\u5728\u9762\u5bf9\u53ef\u7591\u8272\u7d20\u6027\u75c5\u7076\u65f6\uff0c\u5e94\u5982\u4f55\u8bb0\u5f55\u5371\u9669\u4fe1\u53f7\u3001\u5982\u4f55\u89e3\u91ca\u5f53\u524d\u8bc1\u636e\u4e0d\u8db3\u3001\u4f55\u65f6\u5efa\u8bae\u75c5\u7406\u6d3b\u68c0\uff0c\u4ee5\u53ca\u5982\u4f55\u5b89\u6392\u590d\u8bca\u89c2\u5bdf\u3002',
        'tips': '\u56de\u7b54\u60a3\u8005\u95ee\u9898\u65f6\uff0c\u8981\u907f\u514d\u628a\u7ecf\u9a8c\u5224\u65ad\u8bf4\u6210\u6700\u7ec8\u7ed3\u8bba\uff0c\u8981\u660e\u786e\u54ea\u4e9b\u5185\u5bb9\u6765\u81ea\u75c5\u53f2\u3001\u54ea\u4e9b\u6765\u81ea\u76ae\u80a4\u955c\u3001\u54ea\u4e9b\u4ecd\u9700\u75c5\u7406\u786e\u8ba4\u3002'
    },
    {
        'folder': '\u6f14\u793a\u533b\u751f\u5e93-\u6e7f\u75b9\u4e0e\u6162\u6027\u76ae\u708e\u7ba1\u7406\u8d44\u6599',
        'title': '\u6e7f\u75b9\u4e0e\u6162\u6027\u76ae\u708e\u7ba1\u7406\u8d44\u6599',
        'theme': '\u56f4\u7ed5\u6e7f\u75b9\u3001\u7279\u5e94\u6027\u76ae\u708e\u3001\u6162\u6027\u76ae\u708e\u53ca\u590d\u53d1\u7ba1\u7406\u7684\u95e8\u8bca\u5b66\u4e60\u8d44\u6599\u3002',
        'focus': '\u5f3a\u8c03\u8bf1\u56e0\u8ffd\u8e2a\u3001\u76ae\u635f\u5206\u578b\u3001\u5916\u7528\u836f\u89c4\u8303\u3001\u76ae\u80a4\u5c4f\u969c\u4fee\u590d\u3001\u751f\u6d3b\u65b9\u5f0f\u7ba1\u7406\uff0c\u4ee5\u53ca\u590d\u53d1\u671f\u548c\u7f13\u89e3\u671f\u4e0d\u540c\u7684\u5ba3\u6559\u91cd\u70b9\u3002',
        'tips': '\u533b\u751f\u5728\u89e3\u91ca\u75c5\u60c5\u65f6\uff0c\u5e94\u63d0\u9192\u60a3\u8005\u8fd9\u662f\u5bb9\u6613\u53cd\u590d\u7684\u708e\u75c7\u6027\u75be\u75c5\uff0c\u6cbb\u7597\u76ee\u6807\u5f80\u5f80\u662f\u63a7\u5236\u75c7\u72b6\u548c\u5ef6\u957f\u7a33\u5b9a\u671f\uff0c\u800c\u4e0d\u662f\u7b80\u5355\u8ffd\u6c42\u4e00\u6b21\u6027\u6839\u6cbb\u3002'
    },
    {
        'folder': '\u6f14\u793a\u533b\u751f\u5e93-\u94f6\u5c51\u75c5\u957f\u671f\u968f\u8bbf\u4e0e\u6c9f\u901a\u624b\u518c',
        'title': '\u94f6\u5c51\u75c5\u957f\u671f\u968f\u8bbf\u4e0e\u6c9f\u901a\u624b\u518c',
        'theme': '\u56f4\u7ed5\u94f6\u5c51\u75c5\u957f\u671f\u7ba1\u7406\u3001\u7597\u6548\u8bc4\u4f30\u3001\u4e0d\u826f\u53cd\u5e94\u89c2\u5bdf\u4e0e\u60a3\u8005\u957f\u671f\u4f9d\u4ece\u6027\u6574\u7406\u7684\u5b66\u4e60\u8d44\u6599\u3002',
        'focus': '\u5f3a\u8c03\u5206\u5c42\u8bc4\u4f30\u3001\u5173\u8282\u75c7\u72b6\u7b5b\u67e5\u3001\u7cfb\u7edf\u6cbb\u7597\u524d\u6c9f\u901a\u3001\u968f\u8bbf\u8282\u594f\u3001\u7167\u7247\u7559\u5b58\u4e0e\u6cbb\u7597\u76ee\u6807\u8bbe\u5b9a\uff0c\u5e2e\u52a9\u533b\u751f\u5efa\u7acb\u6301\u7eed\u7ba1\u7406\u601d\u8def\u3002',
        'tips': '\u95ee\u7b54\u548c\u5ba3\u6559\u65f6\u8981\u517c\u987e\u4e13\u4e1a\u6027\u548c\u53ef\u6267\u884c\u6027\uff0c\u8ba9\u60a3\u8005\u77e5\u9053\u4e3a\u4ec0\u4e48\u8981\u590d\u8bca\u3001\u4e3a\u4ec0\u4e48\u8981\u76d1\u6d4b\uff0c\u4ee5\u53ca\u54ea\u4e9b\u53d8\u5316\u9700\u8981\u63d0\u524d\u5c31\u533b\u3002'
    },
]

FILES = [MD_NAME, TXT_NAME, CSV_NAME, XLSX_NAME, XLS_NAME, DOCX_NAME, PDF_NAME]


def ensure_clean_dir(path: Path):
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def long_text(spec: dict, label: str) -> str:
    return (
        spec['title'] + '的' + label + '主要用于演示医生知识库的真实内容形态。'
        + '本资料的主题是：' + spec['theme']
        + '在日常门诊中，' + spec['focus']
        + '资料编写时特别强调：所有记录都要尽量写成完整句子，而不是只留下零碎关键词，因为只有完整的上下文，后续做检索和问答时才更容易命中准确语义。'
        + '对于黑色素瘤、湿疹或银屑病这类需要长期判断和随访的场景，医生既要记录当前看到的皮损表现，也要记录患者自述的时间线、既往处理方式和复诊计划。'
        + spec['tips']
        + '此外，这些资料还要提醒医生在形成建议时区分经验性判断、指南性建议和必须进一步检查后才能确认的内容，避免因为表达过于绝对而造成误解。'
        + '本文件统一采用 UTF-8 编码生成，内容明确、完整、可直接阅读，不允许出现乱码、问号占位或无法解释的缩写。'
    )


def markdown_text(spec: dict) -> str:
    return (
        '# ' + spec['title'] + '\n\n'
        + long_text(spec, 'Markdown总览') + '\n\n'
        + '## 门诊使用原则\n'
        + '1. 接诊时先判断当前问题属于初筛、复诊还是治疗后随访。\n'
        + '2. 记录病史时要把主诉、时间线、诱因和既往处理拆开写清楚。\n'
        + '3. 如果证据不足，建议明确写出还需要什么检查，而不是直接下最终结论。\n'
        + '4. 医生沟通时应兼顾专业性与可执行性，让患者知道下一步该做什么。\n\n'
        + '## 文档用途\n'
        + '本文件用于演示知识库上传、解析、切分、检索和问答效果，因此正文使用较完整的中文段落，方便系统进行语义理解。\n'
    )


def txt_text(spec: dict) -> str:
    return (
        spec['title'] + '\n\n'
        + long_text(spec, '纯文本学习摘要') + '\n\n'
        + '补充提醒：医生个人学习资料除了记录疾病知识，还应该记录门诊表达方式、患者常见误区、复诊提示语和风险沟通边界。'
        + '这样在问答系统里检索时，不仅能回答疾病本身的问题，也能回答“怎么向患者解释”这类更贴近临床工作的提问。'
    )


def csv_rows(spec: dict):
    return [
        ['序号', '主题场景', '记录重点', '门诊提醒'],
        ['1', '初诊评估', long_text(spec, '初诊评估部分')[:300], '先记录客观表现，再补充病史和下一步计划'],
        ['2', '复诊随访', long_text(spec, '复诊随访部分')[:300], '强调变化趋势、依从性和复诊节奏'],
        ['3', '患者沟通', long_text(spec, '患者沟通部分')[:300], '避免绝对化表达，说明证据边界'],
        ['4', '风险提示', long_text(spec, '风险提示部分')[:300], '对需要进一步检查的情况要明确写出原因'],
    ]


def write_csv(path: Path, spec: dict):
    with path.open('w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerows(csv_rows(spec))


def write_xlsx(path: Path, spec: dict):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '学习资料'
    for row in csv_rows(spec):
        ws.append(row)
    wb.save(path)


def write_xls(path: Path, spec: dict):
    import xlwt
    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('学习资料')
    for r, row in enumerate(csv_rows(spec)):
        for c, value in enumerate(row):
            ws.write(r, c, value)
    wb.save(str(path))


def write_docx(path: Path, spec: dict):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading(spec['title'] + ' - 医生注意事项', level=1)
    for label in ['接诊前提醒', '问诊记录提醒', '复诊沟通提醒', '长期随访提醒']:
        doc.add_paragraph(long_text(spec, label))
    doc.save(path)


def write_pdf(path: Path, spec: dict):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_candidates = [
        Path(r'C:\\Windows\\Fonts\\msyh.ttc'),
        Path(r'C:\\Windows\\Fonts\\msyh.ttf'),
        Path(r'C:\\Windows\\Fonts\\simhei.ttf'),
    ]
    font_path = None
    for candidate in font_candidates:
        if candidate.exists():
            font_path = candidate
            break
    if font_path is None:
        raise RuntimeError('未找到可用中文字体，无法生成 PDF')

    pdfmetrics.registerFont(TTFont('CNFONT', str(font_path)))
    styles = getSampleStyleSheet()
    body = ParagraphStyle('BodyCN', parent=styles['BodyText'], fontName='CNFONT', fontSize=11, leading=16, spaceAfter=6)
    title = ParagraphStyle('TitleCN', parent=styles['Title'], fontName='CNFONT', fontSize=16, leading=20, spaceAfter=10)
    h2 = ParagraphStyle('H2CN', parent=styles['Heading2'], fontName='CNFONT', fontSize=12, leading=16, spaceAfter=8)

    story = [Paragraph(spec['title'], title), Spacer(1, 4)]
    for heading, label in [
        ('学习背景', '学习背景部分'),
        ('临床记录建议', '临床记录建议部分'),
        ('患者沟通边界', '患者沟通边界部分'),
        ('复诊与随访重点', '复诊与随访重点部分'),
    ]:
        story.append(Paragraph(heading, h2))
        story.append(Paragraph(long_text(spec, label), body))

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    doc.build(story)


def validate_text(text: str, file_label: str):
    if '???' in text:
        raise RuntimeError(file_label + ' 包含 ???')
    if '\ufffd' in text:
        raise RuntimeError(file_label + ' 包含替换字符')
    cleaned = text.replace('\r', '').replace('\n', '').strip()
    if len(cleaned) < 250:
        raise RuntimeError(file_label + ' 正文少于250字，当前 ' + str(len(cleaned)) + ' 字')


ensure_clean_dir(TARGET)
manifest = []

for spec in SUITES:
    folder = TARGET / spec['folder']
    folder.mkdir(parents=True, exist_ok=True)

    md_text = markdown_text(spec)
    txt_text_content = txt_text(spec)

    validate_text(md_text, spec['folder'] + '/' + MD_NAME)
    validate_text(txt_text_content, spec['folder'] + '/' + TXT_NAME)

    (folder / MD_NAME).write_text(md_text, encoding='utf-8')
    (folder / TXT_NAME).write_text(txt_text_content, encoding='utf-8')
    write_csv(folder / CSV_NAME, spec)
    write_xlsx(folder / XLSX_NAME, spec)
    write_xls(folder / XLS_NAME, spec)
    write_docx(folder / DOCX_NAME, spec)
    write_pdf(folder / PDF_NAME, spec)

    manifest.append({'folder': spec['folder'], 'files': FILES})

(TARGET / 'manifest.json').write_text(
    json.dumps({'root': str(TARGET), 'suite_count': len(SUITES), 'items': manifest}, ensure_ascii=False, indent=2),
    encoding='utf-8'
)

print(json.dumps({'status': 'ok', 'root': str(TARGET), 'suite_count': len(SUITES), 'folders': [s['folder'] for s in SUITES]}, ensure_ascii=False, indent=2))
